"""05 - Cek silang spreadsheet: bangun workbook dengan rumus SUMIFS, hitung ulang dengan LibreOffice Calc
(Excel tidak tersedia di sesi), lalu bandingkan dengan hasil SQL.

Output: project/results/2-analysis/q14_excel_crosscheck.xlsx (rumus, sudah dihitung ulang)
        project/results/2-analysis/05_crosscheck_spreadsheet.csv
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from common import CLEAN_CSV, RES_ANALYSIS  # noqa: E402

OUT = RES_ANALYSIS / "q14_excel_crosscheck.xlsx"
tx = pd.read_csv(CLEAN_CSV, parse_dates=["invoice_ts"])
tx["month"] = tx["invoice_ts"].dt.to_period("M").dt.to_timestamp()
agg = (tx.groupby(["month", "country"])["line_revenue"].sum().round(2).reset_index()
       .sort_values(["month", "country"]))

wb = Workbook()
data = wb.active
data.title = "data_bulan_negara"
data.append(["bulan", "negara", "net_revenue"])
for r in agg.itertuples(index=False):
    data.append([r.month.to_pydatetime().date(), r.country, float(r.line_revenue)])
n = len(agg) + 1
for c in data[1]:
    c.font = Font(bold=True)

months = sorted(agg["month"].unique())
cm = wb.create_sheet("cek_bulan")
cm.append(["bulan", "net_revenue_excel"])
for i, m in enumerate(months, start=2):
    cm.append([pd.Timestamp(m).date(), f"=SUMIFS(data_bulan_negara!$C$2:$C${n},data_bulan_negara!$A$2:$A${n},A{i})"])

countries = sorted(agg["country"].unique())
cc = wb.create_sheet("cek_negara")
cc.append(["negara", "net_revenue_excel"])
for i, c in enumerate(countries, start=2):
    cc.append([c, f"=SUMIFS(data_bulan_negara!$C$2:$C${n},data_bulan_negara!$B$2:$B${n},A{i})"])
for ws in (cm, cc):
    for c in ws[1]:
        c.font = Font(bold=True)

# Hitung ulang dengan LibreOffice Calc (headless) agar nilai rumus tersimpan
tmp = Path(tempfile.mkdtemp(prefix="xcheck_"))
src = tmp / "in.xlsx"
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)],
               check=True, capture_output=True, timeout=600)
shutil.copy(tmp / "out" / "in.xlsx", OUT)

# Bandingkan nilai hasil LibreOffice dengan SQL
calc = load_workbook(OUT, data_only=True)
q1 = pd.read_csv(RES_ANALYSIS / "q1_monthly.csv")
q2 = pd.read_csv(RES_ANALYSIS / "q2_country.csv").set_index("country")
rows = []
xl_m = [row[1] for row in calc["cek_bulan"].iter_rows(min_row=2, values_only=True)]
sql_m = q1["net_revenue"].values
d_m = [abs(float(a) - float(b)) for a, b in zip(xl_m, sql_m)]
rows.append({"level": "net revenue per bulan (LibreOffice)", "dibandingkan": len(sql_m),
             "sama": sum(d <= 0.01 for d in d_m), "selisih_maks": round(max(d_m), 6)})
xl_c = {row[0]: row[1] for row in calc["cek_negara"].iter_rows(min_row=2, values_only=True)}
d_c = [abs(float(xl_c[c]) - float(q2.loc[c, "net_revenue"])) for c in q2.index]
rows.append({"level": "net revenue per negara (LibreOffice)", "dibandingkan": len(d_c),
             "sama": sum(d <= 0.01 for d in d_c), "selisih_maks": round(max(d_c), 6)})
out = pd.DataFrame(rows)
out["status"] = ["PASS" if s == d else "FAIL" for s, d in zip(out["sama"], out["dibandingkan"])]
out.to_csv(RES_ANALYSIS / "05_crosscheck_spreadsheet.csv", index=False)
print(out.to_string(index=False))
